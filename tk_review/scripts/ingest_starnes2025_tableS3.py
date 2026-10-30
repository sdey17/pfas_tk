#!/usr/bin/env python3
"""Starnes et al. 2025 Table S3 -- chemical stock preparation details.

WHY THIS EXISTS, AND WHAT IT IS NOT. The database carried 68 Starnes
2025 rows whose ligand concentration read "concentration-response series
per chemical (Table S3, not extracted)". That note was wrong about what
Table S3 holds. It was fetched and read, and Table S3 is titled
"Chemical stock preparation details": the STOCK concentration (10 or
20 mM) and the solvent each chemical was dissolved in. The
concentration-response series is not in it.

Nor is it anywhere else in the paper. The methods say only that stocks
"were serially diluted to achieve necessary working stock concentrations
for DSF concentration-response analysis" and refer the design to Jackson
et al. 2021 and Starnes et al. 2024. No table, figure or sentence in the
paper or its seven supplementary tables gives the concentrations tested.
So those 68 rows are not an extraction gap; the primary source does not
report the number, and the note now says so.

WHAT IS WORTH HAVING, and is why this file is kept. Six chemicals could
not be dissolved in buffer alone and were prepared in 20% to 50% DMSO:

    50% DMSO   9H-perfluorononanoic acid, 6:2 FTCA, EtFOSA, EtFOSE
    25% DMSO   2-(perfluorobutyl)ethyl acrylate
    20% DMSO   Nonafluoropentanamide

That is a co-solvent at a concentration which is not a minor detail for
an assay that reads out albumin's melting temperature: DMSO destabilises
albumin and competes for its hydrophobic sites, and the measurement is a
thermal shift. Those six rows are not on the same footing as the other
62 and the database could not previously tell them apart.

SOURCE. doi 10.1093/toxsci/kfae124, PMC11664106, supplementary workbook
toxsci-24-0213-File008.xlsx, sheet "Table S3", fetched from the Europe
PMC supplementary-files endpoint for PMC11664106. The workbook is 135 MB
(its machine-learning feature matrices) and is NOT kept in the
repository; this script re-derives the CSV from it if it is fetched
again.

Run:  python3 -I tk_review/scripts/ingest_starnes2025_tableS3.py <xlsx>
"""
import csv
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "db", "starnes2025_tableS3_stocks.csv")

FIELDS = ["cas_rn", "chemical_name", "abbreviations", "stock_conc_mM",
          "solvent", "cosolvent_percent", "source_table"]
SOURCE = ("Starnes et al. 2025, Toxicol Sci 203:67, doi "
          "10.1093/toxsci/kfae124, supplementary Table S3")
PCT = re.compile(r"(\d+)\s*%")


def rows(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Table S3"]
    out, header_seen = [], False
    for row in ws.iter_rows(values_only=True):
        cells = ["" if c is None else str(c).strip() for c in row]
        # The sheet is indented one column and carries a title row, a
        # header row and a footnote row; key off the header instead of a
        # fixed offset so a re-export that shifts does not go unnoticed.
        if not header_seen:
            header_seen = any(c == "CAS-RN" for c in cells)
            continue
        cas = next((c for c in cells if re.fullmatch(r"\d{2,7}-\d\d-\d", c)),
                   "")
        if not cas:
            continue                      # the footnote row
        i = cells.index(cas)
        name, abbr, conc, solvent = (cells[i + 1:i + 5] + ["", "", "", ""])[:4]
        m = PCT.search(solvent)
        out.append({
            "cas_rn": cas, "chemical_name": name, "abbreviations": abbr,
            "stock_conc_mM": conc, "solvent": solvent,
            "cosolvent_percent": m.group(1) if m else "",
            "source_table": SOURCE})
    wb.close()
    return out


def main():
    if len(sys.argv) < 2:
        print(__doc__.split("Run:")[0].rstrip())
        print("\nusage: ingest_starnes2025_tableS3.py <supplementary.xlsx>")
        return 2
    recs = rows(sys.argv[1])
    with open(OUT, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(recs)
    print(f"{len(recs)} chemicals from Table S3")
    conc = sorted({r["stock_conc_mM"] for r in recs})
    print(f"  stock concentrations: {', '.join(conc)} mM")
    co = [r for r in recs if r["cosolvent_percent"]]
    print(f"  {len(co)} prepared with a DMSO co-solvent:")
    for r in sorted(co, key=lambda r: -int(r["cosolvent_percent"])):
        print(f"    {r['cosolvent_percent']:>3s}%  {r['cas_rn']:14s} "
              f"{(r['abbreviations'] or r['chemical_name'])[:52]}")
    print(f"  {len(recs) - len(co)} in buffer alone")
    print(f"\nwrote {os.path.abspath(OUT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
